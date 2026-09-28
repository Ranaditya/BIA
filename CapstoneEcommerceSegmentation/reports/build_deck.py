"""Build the capstone presentation from the BIA template and the pipeline outputs.

Usage, from the repository root after running ``main.py``::

	python CapstoneEcommerceSegmentation/reports/build_deck.py

Charts, tables, and most figures are read from ``outputs/dashboard_data.json``
and ``outputs/analysis_results.json``. Some narrative text (skewness values,
RFM medians, k=4 overlap counts, and speaker-note talking points) reflects the
recorded k=2 run and must be reviewed if the data or pipeline changes.

``dashboard_screenshot.png`` is a headless-browser capture of
``dashboards/index.html``; recapture it if the dashboard changes.
Requires ``python-pptx``.
"""

import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

REPORTS = Path(__file__).resolve().parent
PROJECT = REPORTS.parent
TEMPLATE = REPORTS / "Capstone PPT Template.pptx"
OUT = REPORTS / "Capstone_Ecommerce_Segmentation.pptx"
DASH_PNG = REPORTS / "dashboard_screenshot.png"

dash = json.loads((PROJECT / "outputs" / "dashboard_data.json").read_text())
res = json.loads((PROJECT / "outputs" / "analysis_results.json").read_text())
sig = dash["signature"]

# ---------- palette (from template navy + network-glow blue) ----------
NAVY = RGBColor(0x1D, 0x1B, 0x58)
BLUE = RGBColor(0x1B, 0x8F, 0xE0)
SKY = RGBColor(0x7F, 0xB8, 0xEA)
ICE = RGBColor(0xEC, 0xF1, 0xF9)
LINE = RGBColor(0xD5, 0xDB, 0xE8)
INK = RGBColor(0x22, 0x24, 0x33)
MUTED = RGBColor(0x5B, 0x60, 0x75)
GREY = RGBColor(0xA9, 0xB1, 0xC6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RISK = RGBColor(0xE4, 0x57, 0x2E)
GAIN = RGBColor(0x1E, 0x9E, 0x6A)
ON_NAVY_MUTED = RGBColor(0xB9, 0xC2, 0xE6)
SEG = {"High-Value Active": BLUE, "Low-Value Lapsing": GREY,
	   "Mid-Value Active": SKY, "Mid-Value Cooling": RGBColor(0x6C, 0x74, 0x99)}
FONT = "Calibri"

prs = Presentation(TEMPLATE)
L_TITLE_CONTENT, L_ONE_THIRD, L_ONE_THIRD_R, L_HALF = 1, 2, 3, 4
L_CHEVRON, L_SECTION, L_TITLE_ONLY = 7, 8, 11

# ---------- primitives ----------

def box(slide, x, y, w, h, text="", size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
		anchor=MSO_ANCHOR.TOP, italic=False, font=FONT):
	tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
	tf = tb.text_frame
	tf.word_wrap = True
	tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
	tf.vertical_anchor = anchor
	lines = text if isinstance(text, list) else [text]
	for i, line in enumerate(lines):
		p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
		p.alignment = align
		runs = line if isinstance(line, list) else [(line, {})] if isinstance(line, str) else [line]
		for run_text, fmt in runs:
			r = p.add_run()
			r.text = run_text
			f = r.font
			f.name = font
			f.size = Pt(fmt.get("size", size))
			f.bold = fmt.get("bold", bold)
			f.italic = fmt.get("italic", italic)
			f.color.rgb = fmt.get("color", color)
	return tb


def bullets(slide, x, y, w, h, items, size=14, color=INK, gap=6, bullet_color=BLUE):
	"""items: list of str or list of (text, fmt) runs; tuple ('sub', text) for level 2."""
	tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
	tf = tb.text_frame
	tf.word_wrap = True
	tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
	for i, item in enumerate(items):
		level = 0
		if isinstance(item, tuple) and item[0] == "sub":
			level, item = 1, item[1]
		p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
		p.space_after = Pt(gap)
		runs = item if isinstance(item, list) else [(item, {})]
		for run_text, fmt in runs:
			r = p.add_run()
			r.text = run_text
			r.font.name = FONT
			r.font.size = Pt(fmt.get("size", size - 2 * level))
			r.font.bold = fmt.get("bold", False)
			r.font.color.rgb = fmt.get("color", color if level == 0 else MUTED)
		pPr = p._p.get_or_add_pPr()
		indent = 228600 if level == 0 else 457200
		pPr.set("marL", str(indent))
		pPr.set("indent", str(-228600))
		buClr = etree.SubElement(pPr, qn("a:buClr"))
		etree.SubElement(buClr, qn("a:srgbClr")).set("val", str(bullet_color))
		etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
		etree.SubElement(pPr, qn("a:buChar")).set("char", "•" if level == 0 else "–")
	return tb


def rect(slide, x, y, w, h, fill=ICE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
	s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
	if fill is None:
		s.fill.background()
	else:
		s.fill.solid()
		s.fill.fore_color.rgb = fill
	if line is None:
		s.line.fill.background()
	else:
		s.line.color.rgb = line
		s.line.width = Pt(1)
	if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
		s.adjustments[0] = radius
	s.shadow.inherit = False
	s.text_frame.text = ""
	return s


def badge(slide, x, y, d, text, fill=NAVY, color=WHITE, size=14):
	c = rect(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
	tf = c.text_frame
	tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
	tf.vertical_anchor = MSO_ANCHOR.MIDDLE
	p = tf.paragraphs[0]
	p.alignment = PP_ALIGN.CENTER
	r = p.add_run()
	r.text = text
	r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), True, color
	return c


def stat(slide, x, y, w, value, label, vcolor=NAVY, lcolor=MUTED, vsize=40, lsize=12):
	box(slide, x, y, w, vsize / 72 * 1.15, value, size=vsize, color=vcolor, bold=True)
	box(slide, x, y + vsize / 72 * 1.2, w, 0.6, label, size=lsize, color=lcolor)


def new_slide(layout, title=None):
	s = prs.slides.add_slide(prs.slide_layouts[layout])
	for ph in list(s.placeholders):
		if ph.placeholder_format.type == 1 and title is not None:  # TITLE
			ph.text_frame.text = title
		else:
			ph._element.getparent().remove(ph._element)
	return s


# ---------- chart styling ----------

def style_axes(chart, color=MUTED, grid=LINE, size=11, value_fmt=None, show_val_axis=True):
	chart.font.name = FONT
	chart.font.size = Pt(size)
	chart.font.color.rgb = color
	try:
		va = chart.value_axis
		va.has_major_gridlines = grid is not None
		if grid is not None:
			va.major_gridlines.format.line.color.rgb = grid
			va.major_gridlines.format.line.width = Pt(0.75)
		va.format.line.fill.background()
		va.tick_labels.font.size = Pt(size)
		va.tick_labels.font.color.rgb = color
		if value_fmt:
			va.tick_labels.number_format = value_fmt
			va.tick_labels.number_format_is_linked = False
		va.visible = show_val_axis
		ca = chart.category_axis
		ca.format.line.color.rgb = grid or color
		ca.tick_labels.font.size = Pt(size)
		ca.tick_labels.font.color.rgb = color
		ca.has_major_gridlines = False
	except (ValueError, AttributeError):
		pass


def labels_on(plot, fmt, size=11, color=INK, pos=XL_LABEL_POSITION.OUTSIDE_END, bold=False):
	plot.has_data_labels = True
	dl = plot.data_labels
	dl.number_format = fmt
	dl.number_format_is_linked = False
	dl.font.size = Pt(size)
	dl.font.color.rgb = color
	dl.font.bold = bold
	if pos is not None:
		dl.position = pos


def color_points(series, colors):
	for i, c in enumerate(colors):
		pt = series.points[i]
		pt.format.fill.solid()
		pt.format.fill.fore_color.rgb = c


def chart_title(chart, text, size=13, color=INK):
	chart.has_title = True
	tf = chart.chart_title.text_frame
	tf.text = text
	r = tf.paragraphs[0].runs[0]
	r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, color, FONT


def axis_title(axis, text, size=11, color=MUTED):
	axis.has_title = True
	tf = axis.axis_title.text_frame
	tf.text = text
	r = tf.paragraphs[0].runs[0]
	r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), False, color, FONT


# ---------- data ----------
seg = {s["Segment"]: s for s in res["segment_profile"]}
hva, lvl = seg["High-Value Active"], seg["Low-Value Lapsing"]
n_cust = sig["n_customers"]
hva_share = hva["Size"] / n_cust * 100
lvl_share = lvl["Size"] / n_cust * 100
hva_rev = hva["Size"] * hva["Monetary"]
lvl_rev = lvl["Size"] * lvl["Monetary"]
lvl_aov = lvl["Monetary"] / lvl["Frequency"]
winback_n = round(lvl["Size"] * 0.10)
winback_gain = winback_n * lvl_aov
risk_n = hva["Size"] * 0.05
risk_loss = risk_n * hva["Monetary"]
cm = dash["confusion"]["matrix"]
test_n = sum(map(sum, cm))
correct = cm[0][0] + cm[1][1]
acc = res["classifier_metrics"]["accuracy"]
stab = res["stability"]
cmp2 = {r["algorithm"]: r for r in res["algorithm_comparison"]}
cmp4 = {r["algorithm"]: r for r in res["alternative_comparison"]}
ksel = dash["k_selection"]
fi = {r["feature"]: r["importance"] for r in res["feature_importance"]}
cent = res["scaled_centroids"]
gap = {f: abs(cent["0"][f] - cent["1"][f]) for f in ["Recency", "Frequency", "Monetary"]}
steps = res["cleaning_steps"]
uk_share = dash["countries"]["transactions"][0] / sig["clean_rows"] * 100
total_rev = hva_rev + lvl_rev
uk_rev_share = dash["countries"]["revenue"][0] / total_rev * 100
months = dash["monthly"]
peak_i = max(range(len(months["revenue"])), key=lambda i: months["revenue"][i])
MONTH_LABEL = {"01": "Jan", "02": "Feb", "03": "Mar", "04": "Apr", "05": "May", "06": "Jun",
			   "07": "Jul", "08": "Aug", "09": "Sep", "10": "Oct", "11": "Nov", "12": "Dec"}
mlabels = [f"{MONTH_LABEL[m[5:]]} '{m[2:4]}" for m in months["labels"]]


def gbp(v, k=False):
	if v >= 1e6:
		return f"£{v / 1e6:.2f}M"
	if k or v >= 1e4:
		return f"£{v / 1e3:.0f}k"
	return f"£{v:,.0f}"


# =====================================================================
# 1. Title (template slide 1)
s_title = prs.slides[0]
for sh in s_title.shapes:
	if sh.has_text_frame and "Capstone" in sh.text_frame.text:
		runs = sh.text_frame.paragraphs[0].runs
		runs[0].text = "Who Drives the Revenue?"
		for r in runs[1:]:
			r.text = ""
box(s_title, 1.0, 4.05, 11.33, 0.5, "RFM-Based Customer Segmentation for a UK Online Retailer",
	size=22, color=WHITE, align=PP_ALIGN.CENTER)
box(s_title, 1.0, 4.75, 11.33, 0.4, "Capstone Project  ·  [Presenter name]  ·  [Date]",
	size=14, color=ON_NAVY_MUTED, align=PP_ALIGN.CENTER)

# 2. Agenda (template slide 2)
s_agenda = prs.slides[1]
agenda = [
	("01", "Problem & data", "Business question, dataset, cleaning"),
	("02", "Customer behaviour", "EDA and RFM feature engineering"),
	("03", "Building segments", "Choosing k and comparing algorithms"),
	("04", "Segments & drivers", "Who they are and what separates them"),
	("05", "Acting on it", "Scoring, recommendations, dashboard, next steps"),
]
cw, gapx, x0, y0 = 2.18, 0.24, 0.74, 2.05
for i, (num, head, sub) in enumerate(agenda):
	x = x0 + i * (cw + gapx)
	rect(s_agenda, x, y0, cw, 3.3, fill=ICE)
	badge(s_agenda, x + 0.3, y0 + 0.35, 0.72, num, size=18)
	box(s_agenda, x + 0.3, y0 + 1.35, cw - 0.55, 0.8, head, size=18, color=NAVY, bold=True)
	box(s_agenda, x + 0.3, y0 + 2.2, cw - 0.55, 0.9, sub, size=13, color=MUTED)
box(s_agenda, 0.74, 5.62, 11.8, 0.3,
	"Backup: an appendix of anticipated methodology questions (FAQ 1–9) follows the close.",
	size=12, color=MUTED, italic=True)

# 3. Business problem
s = new_slide(L_ONE_THIRD, "The business problem")
box(s, 5.1, 0.85, 7.4, 1.0, "“Which customers generate our revenue, which are we losing, and how should "
	"marketing treat each group differently?”", size=20, color=NAVY, italic=True)
objs = [
	("Segment", "Group customers by purchasing behaviour (Recency, Frequency, Monetary) and compare "
	 "K-means, hierarchical clustering and DBSCAN."),
	("Explain", "Show which attributes drive the segments and how much revenue each segment holds."),
	("Operationalise", "Score new customers into a segment with a classifier and surface results in an "
	 "interactive dashboard."),
	("Act", "Recommend targeted retention and marketing strategies, sized and testable."),
]
for i, (h, t) in enumerate(objs):
	y = 2.2 + i * 0.95
	badge(s, 5.1, y, 0.55, str(i + 1), fill=BLUE, size=16)
	box(s, 5.9, y - 0.02, 6.6, 0.35, h, size=16, color=NAVY, bold=True)
	box(s, 5.9, y + 0.3, 6.6, 0.6, t, size=13, color=MUTED)

# 4. Data & cleaning
s = new_slide(L_TITLE_ONLY, "Data: one year of UK online retail transactions")
rect(s, 0.74, 1.6, 3.9, 4.35, fill=NAVY)
box(s, 1.05, 1.85, 3.3, 0.35, "SOURCE", size=11, color=ON_NAVY_MUTED, bold=True)
box(s, 1.05, 2.15, 3.3, 0.9, "UCI Online Retail dataset", size=20, color=WHITE, bold=True)
box(s, 1.05, 2.85, 3.3, 1.2, "UK-based online gift retailer; many customers are wholesalers.\n"
	"1 Dec 2010 – 9 Dec 2011 · 8 columns", size=13, color=ON_NAVY_MUTED)
stat(s, 1.05, 4.05, 3.3, f"{n_cust:,}", "customers after cleaning", vcolor=WHITE,
	 lcolor=ON_NAVY_MUTED, vsize=36)
cd = CategoryChartData()
cd.categories = ["Raw rows"] + [st["step"] for st in steps]
cd.add_series("Rows remaining", [sig["raw_rows"]] + [st["remaining"] for st in steps])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(5.0), Inches(1.5), Inches(7.6), Inches(4.0), cd)
ch = gf.chart
ch.has_legend = False
chart_title(ch, "Rows remaining after each cleaning step")
style_axes(ch, grid=None, show_val_axis=False)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 45
labels_on(ch.plots[0], "0", size=11)
color_points(ch.plots[0].series[0], [GREY] + [SKY] * (len(steps) - 1) + [NAVY])
for i, v in enumerate([sig["raw_rows"]] + [st["remaining"] for st in steps]):
	dl = ch.plots[0].series[0].points[i].data_label
	dl.text_frame.text = f"{v:,}"
	dl.position = XL_LABEL_POSITION.OUTSIDE_END
	run = dl.text_frame.paragraphs[0].runs[0]
	run.font.size, run.font.color.rgb = Pt(11), INK
box(s, 5.0, 5.55, 7.6, 0.4, f"{sig['raw_rows']:,} → {sig['clean_rows']:,} rows ({sig['clean_rows'] / sig['raw_rows'] * 100:.0f}% kept). "
	f"The biggest cut, missing CustomerID ({steps[0]['removed']:,} rows), is unavoidable: RFM needs a customer.",
	size=12, color=MUTED)

# 5. EDA
s = new_slide(L_TITLE_ONLY, "What the transactions tell us")
cd = CategoryChartData()
cd.categories = mlabels
cd.add_series("Revenue", months["revenue"])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.74), Inches(1.5), Inches(7.5), Inches(4.1), cd)
ch = gf.chart
ch.has_legend = False
chart_title(ch, "Monthly revenue (£)")
style_axes(ch, value_fmt='£#,##0,"k"')
ser = ch.plots[0].series[0]
ser.smooth = False
ser.format.line.color.rgb = NAVY
ser.format.line.width = Pt(2.5)
ser.marker.style = XL_MARKER_STYLE.CIRCLE
ser.marker.size = 7
ser.marker.format.fill.solid()
ser.marker.format.fill.fore_color.rgb = NAVY
ser.marker.format.line.color.rgb = NAVY
pk = ser.points[peak_i]
pk.marker.style = XL_MARKER_STYLE.CIRCLE
pk.marker.format.fill.solid()
pk.marker.format.fill.fore_color.rgb = BLUE
pk.marker.format.line.color.rgb = BLUE
pk.marker.size = 11
pk.data_label.has_text_frame = True
pk.data_label.text_frame.text = f"Peak {gbp(months['revenue'][peak_i])}"
pk.data_label.position = XL_LABEL_POSITION.LEFT
pk.data_label.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
pk.data_label.text_frame.paragraphs[0].runs[0].font.bold = True
pk.data_label.text_frame.paragraphs[0].runs[0].font.color.rgb = BLUE
box(s, 0.9, 5.7, 7.2, 0.3, "Dec '11 covers 1–9 December only.", size=11, color=MUTED, italic=True)
cards = [
	(f"{uk_share:.0f}%", "of transactions come from the UK", f"({uk_rev_share:.0f}% of revenue)"),
	("Q4 peak", "Sep–Nov 2011 revenue climbs to", f"{gbp(months['revenue'][peak_i])} in Nov: seasonal gifting"),
	("£614 vs £1,853", "median vs mean spend per customer", "a long tail of very large buyers (max £265k)"),
]
for i, (v, l1, l2) in enumerate(cards):
	y = 1.55 + i * 1.5
	rect(s, 8.6, y, 4.0, 1.32, fill=ICE)
	box(s, 8.85, y + 0.12, 3.6, 0.5, v, size=24, color=NAVY, bold=True)
	box(s, 8.85, y + 0.62, 3.6, 0.65, [l1, [(l2, {"color": MUTED, "size": 12})]], size=12, color=INK)

# 6. RFM feature engineering (chevron layout: navy on right)
s = new_slide(L_CHEVRON, "RFM feature engineering")
rfm_cards = [
	("R", "Recency", "Days since last purchase (snapshot 10 Dec 2011)", "median 51 days"),
	("F", "Frequency", "Number of distinct invoices", "median 2 orders"),
	("M", "Monetary", "Total spend = Σ Quantity × UnitPrice", "median £614"),
]
for i, (letter, name, desc, med) in enumerate(rfm_cards):
	y = 1.65 + i * 1.4
	rect(s, 0.74, y, 6.9, 1.2, fill=ICE)
	badge(s, 0.98, y + 0.25, 0.7, letter, fill=NAVY, size=22)
	box(s, 1.95, y + 0.15, 3.3, 0.4, name, size=18, color=NAVY, bold=True)
	box(s, 1.95, y + 0.55, 3.6, 0.6, desc, size=13, color=MUTED)
	box(s, 5.45, y + 0.3, 2.0, 0.6, med, size=15, color=BLUE, bold=True, align=PP_ALIGN.RIGHT,
		anchor=MSO_ANCHOR.MIDDLE)
box(s, 9.55, 1.75, 3.2, 0.4, "THEN TRANSFORM", size=11, color=ON_NAVY_MUTED, bold=True)
box(s, 9.55, 2.1, 3.2, 1.0, ["log1p", "→ StandardScaler"], size=22, color=WHITE, bold=True)
box(s, 9.55, 3.2, 3.1, 2.3, [
	"Monetary skewness falls from 20.4 to 0.4; Frequency from 11.8 to 1.2.",
	"",
	"Scaling puts R, F and M on equal footing for distance-based clustering."],
	size=13, color=ON_NAVY_MUTED)

# 7. Choosing k
s = new_slide(L_TITLE_ONLY, "Choosing k: the silhouette favours two segments")
ks = [str(r["k"]) for r in ksel]
cd = CategoryChartData()
cd.categories = ks
cd.add_series("Silhouette", [r["silhouette"] for r in ksel])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.74), Inches(1.5), Inches(5.8), Inches(3.9), cd)
ch = gf.chart
ch.has_legend = False
chart_title(ch, "Silhouette score by k (higher is better)")
style_axes(ch, grid=None, show_val_axis=False)
ch.plots[0].gap_width = 50
labels_on(ch.plots[0], "0.00", size=10)
color_points(ch.plots[0].series[0], [BLUE] + [GREY] * (len(ks) - 1))
axis_title(ch.category_axis, "k (number of clusters)")
cd = CategoryChartData()
cd.categories = ks
cd.add_series("Inertia", [r["inertia"] for r in ksel])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(6.8), Inches(1.5), Inches(5.8), Inches(3.9), cd)
ch = gf.chart
ch.has_legend = False
chart_title(ch, "Inertia by k (elbow method)")
style_axes(ch, value_fmt="#,##0")
ser = ch.plots[0].series[0]
ser.smooth = False
ser.format.line.color.rgb = NAVY
ser.format.line.width = Pt(2.25)
ser.marker.style = XL_MARKER_STYLE.CIRCLE
ser.marker.format.fill.solid()
ser.marker.format.fill.fore_color.rgb = NAVY
axis_title(ch.category_axis, "k (number of clusters)")
rect(s, 0.74, 5.45, 11.85, 0.55, fill=ICE)
box(s, 0.95, 5.45, 11.5, 0.55, [[("k = 2 wins clearly: ", {"bold": True, "color": NAVY}),
	(f"silhouette 0.43 vs 0.34 at k = 3. The elbow has no sharp bend, so the silhouette decides. "
	 "k = 4 is a finer view (FAQ 3).", {})]], size=13, anchor=MSO_ANCHOR.MIDDLE)

# 8. Algorithm comparison
s = new_slide(L_TITLE_ONLY, "Comparing algorithms: K-means separates best")
algos = ["KMeans", "Hierarchical", "DBSCAN"]
names = {"KMeans": "K-means", "Hierarchical": "Hierarchical (Ward)", "DBSCAN": "DBSCAN"}
cd = CategoryChartData()
cd.categories = [names[a] for a in algos]
cd.add_series("Silhouette", [cmp2[a]["silhouette"] for a in algos])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.74), Inches(1.5), Inches(5.3), Inches(2.6), cd)
ch = gf.chart
ch.has_legend = False
chart_title(ch, "Silhouette (higher is better)")
style_axes(ch, grid=None, show_val_axis=False)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 45
labels_on(ch.plots[0], "0.000")
color_points(ch.plots[0].series[0], [BLUE, GREY, GREY])
cd = CategoryChartData()
cd.categories = [names[a] for a in algos]
cd.add_series("Davies-Bouldin", [cmp2[a]["davies_bouldin"] for a in algos])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.74), Inches(4.1), Inches(5.3), Inches(1.95), cd)
ch = gf.chart
ch.has_legend = False
chart_title(ch, "Davies-Bouldin (lower is better)")
style_axes(ch, grid=None, show_val_axis=False)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 45
labels_on(ch.plots[0], "0.00")
color_points(ch.plots[0].series[0], [BLUE, GREY, GREY])
rows = [["", "K-means", "Hierarchical", "DBSCAN"],
		["How it groups", "Nearest centroid", "Merge closest (Ward)", "Dense regions"],
		["Clusters found", "2", "2", f"2 + {cmp2['DBSCAN']['noise_points']} noise"],
		["Largest cluster", f"{cmp2['KMeans']['largest_cluster_pct']:.0f}%", f"{cmp2['Hierarchical']['largest_cluster_pct']:.0f}%",
		 f"{cmp2['DBSCAN']['largest_cluster_pct']:.0f}%"],
		["Scores new customers?", "Yes (predict)", "No (refit)", "No (refit)"],
		["Silhouette at k = 4", f"{cmp4['KMeans']['silhouette']:.2f}", f"{cmp4['Hierarchical']['silhouette']:.2f}", "n/a (no k)"]]
tbl = s.shapes.add_table(len(rows), 4, Inches(6.4), Inches(1.6), Inches(6.2), Inches(3.3)).table
tbl.columns[0].width = Inches(1.9)
for c in range(1, 4):
	tbl.columns[c].width = Inches(1.43)
for r_i, row in enumerate(rows):
	for c_i, val in enumerate(row):
		cell = tbl.cell(r_i, c_i)
		cell.text = val or " "
		cell.margin_left = cell.margin_right = Inches(0.08)
		cell.vertical_anchor = MSO_ANCHOR.MIDDLE
		run = cell.text_frame.paragraphs[0].runs[0]
		run.font.name, run.font.size = FONT, Pt(12)
		cell.fill.solid()
		if r_i == 0:
			cell.fill.fore_color.rgb = NAVY
			run.font.color.rgb, run.font.bold = WHITE, True
		else:
			cell.fill.fore_color.rgb = ICE if r_i % 2 else WHITE
			run.font.color.rgb = INK
			run.font.bold = c_i == 0 or c_i == 1
			if c_i == 1:
				run.font.color.rgb = NAVY
box(s, 6.4, 5.1, 6.2, 0.9, "K-means wins on both metrics at k = 2 and k = 4. It is also the only one of the three "
	"that can assign a new customer without refitting.", size=13, color=MUTED)

# 9. Segments + Pareto (Half & Half: navy right half)
s = new_slide(L_HALF, "Two segments, very unequal value")
for i, (name, row, col) in enumerate([("High-Value Active", hva, BLUE), ("Low-Value Lapsing", lvl, GREY)]):
	y = 1.6 + i * 2.2
	rect(s, 0.74, y, 5.1, 2.0, fill=ICE)
	dot = rect(s, 0.98, y + 0.26, 0.26, 0.26, fill=col, shape=MSO_SHAPE.OVAL)
	box(s, 1.38, y + 0.18, 4.3, 0.4, name, size=18, color=NAVY, bold=True)
	box(s, 1.38, y + 0.58, 4.3, 0.3, f"{row['Size']:,} customers · {row['Size'] / n_cust * 100:.0f}%",
		size=12, color=MUTED)
	vals = [(f"{row['Recency']:.0f} d", "last purchase"), (f"{row['Frequency']:.1f}", "orders"),
			(gbp(row["Monetary"]), "avg. spend")]
	for j, (v, lab) in enumerate(vals):
		box(s, 0.98 + j * 1.6, y + 1.0, 1.55, 0.45, v, size=22, color=NAVY, bold=True)
		box(s, 0.98 + j * 1.6, y + 1.47, 1.55, 0.3, lab, size=11, color=MUTED)
cd = CategoryChartData()
cd.categories = ["Customers", "Revenue"]
cd.add_series("High-Value Active", [hva_share / 100, hva["RevenuePct"] / 100])
cd.add_series("Low-Value Lapsing", [lvl_share / 100, lvl["RevenuePct"] / 100])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED_100, Inches(6.75), Inches(1.95), Inches(3.5), Inches(3.9), cd)
ch = gf.chart
style_axes(ch, color=WHITE, grid=None, size=12, show_val_axis=False)
ch.category_axis.format.line.color.rgb = ON_NAVY_MUTED
ch.has_legend = False
ch.plots[0].gap_width = 40
ch.plots[0].overlap = 100
for ser, col, fcol in [(ch.plots[0].series[0], RGBColor(0x4F, 0xB3, 0xFF), NAVY),
					   (ch.plots[0].series[1], RGBColor(0x8C, 0x93, 0xB8), WHITE)]:
	ser.format.fill.solid()
	ser.format.fill.fore_color.rgb = col
	ser.data_labels.show_value = True
	ser.data_labels.number_format = "0%"
	ser.data_labels.number_format_is_linked = False
	ser.data_labels.font.size = Pt(14)
	ser.data_labels.font.bold = True
	ser.data_labels.font.color.rgb = fcol
	ser.data_labels.position = XL_LABEL_POSITION.CENTER
box(s, 6.95, 0.75, 5.0, 0.4, "THE PARETO PATTERN", size=11, color=ON_NAVY_MUTED, bold=True)
box(s, 6.95, 1.1, 5.4, 0.6, f"{hva_share:.0f}% of customers → {hva['RevenuePct']:.0f}% of revenue",
	size=22, color=WHITE, bold=True)
box(s, 10.45, 2.3, 1.6, 2.6, [[("■ ", {"color": RGBColor(0x4F, 0xB3, 0xFF)}), ("High-Value Active", {})], "",
	[("■ ", {"color": RGBColor(0x8C, 0x93, 0xB8)}), ("Low-Value Lapsing", {})]], size=12, color=WHITE)

# 10. What drives the segments
s = new_slide(L_TITLE_ONLY, "What drives the segments? Frequency first")
feats = ["Frequency", "Monetary", "Recency"]
cd = CategoryChartData()
cd.categories = feats
cd.add_series("Gap", [gap[f] for f in feats])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.74), Inches(1.5), Inches(5.7), Inches(3.3), cd)
ch = gf.chart
ch.has_legend = False
chart_title(ch, "Centroid gap between segments (std. units)")
style_axes(ch, grid=None, show_val_axis=False)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 45
labels_on(ch.plots[0], "0.00")
color_points(ch.plots[0].series[0], [BLUE, SKY, GREY])
cd = CategoryChartData()
cd.categories = feats
cd.add_series("Importance", [fi[f] for f in feats])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(6.85), Inches(1.5), Inches(5.7), Inches(3.3), cd)
ch = gf.chart
ch.has_legend = False
chart_title(ch, "Random Forest feature importance")
style_axes(ch, grid=None, show_val_axis=False)
ch.category_axis.reverse_order = True
ch.plots[0].gap_width = 45
labels_on(ch.plots[0], "0%")
color_points(ch.plots[0].series[0], [BLUE, SKY, GREY])
for i, (h, t) in enumerate([
		("Two views agree", "Frequency separates the segments most, then Monetary, then Recency."),
		("F and M move together", "Correlation 0.53: repeat buyers are also big spenders (FAQ 9)."),
		("So what", "Getting a second and third order is the lever that moves a customer up.")]):
	x = 0.74 + i * 4.0
	rect(s, x, 5.0, 3.8, 1.0, fill=ICE)
	box(s, x + 0.2, 5.08, 3.4, 0.3, h, size=13, color=NAVY, bold=True)
	box(s, x + 0.2, 5.4, 3.45, 0.6, t, size=12, color=MUTED)

# 11. Scoring new customers (right navy panel)
s = new_slide(L_ONE_THIRD_R, "Scoring new customers into segments")
flow = ["New customer transactions", "Compute RFM", "log1p + saved scaler", "Random Forest", "Segment label"]
for i, t in enumerate(flow):
	y = 1.6 + i * 0.62
	fill = NAVY if i == len(flow) - 1 else ICE
	rect(s, 0.74, y, 3.1, 0.48, fill=fill)
	box(s, 0.9, y, 2.8, 0.48, t, size=13, color=WHITE if fill == NAVY else NAVY, bold=True,
		anchor=MSO_ANCHOR.MIDDLE)
	if i < len(flow) - 1:
		a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(2.2), Inches(y + 0.49), Inches(0.18), Inches(0.12))
		a.fill.solid()
		a.fill.fore_color.rgb = GREY
		a.line.fill.background()
box(s, 4.35, 1.55, 3.9, 0.35, f"Held-out test set ({test_n} customers)", size=13, color=NAVY, bold=True)
box(s, 5.45, 1.95, 2.8, 0.3, "Predicted", size=11, color=MUTED, align=PP_ALIGN.CENTER)
labs = ["High-Value", "Low-Value"]
for j, lab in enumerate(labs):
	box(s, 5.45 + j * 1.4, 2.25, 1.4, 0.3, lab, size=11, color=MUTED, align=PP_ALIGN.CENTER)
	box(s, 4.35, 2.6 + j * 1.25, 1.0, 1.2, lab, size=11, color=MUTED, align=PP_ALIGN.RIGHT,
		anchor=MSO_ANCHOR.MIDDLE)
for r_i in range(2):
	for c_i in range(2):
		v = cm[r_i][c_i]
		diag = r_i == c_i
		cell = rect(s, 5.45 + c_i * 1.4, 2.6 + r_i * 1.25, 1.35, 1.2, fill=NAVY if diag else ICE, radius=0.04)
		box(s, 5.45 + c_i * 1.4, 2.6 + r_i * 1.25, 1.35, 1.2, f"{v}", size=24, bold=True,
			color=WHITE if diag else RISK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
box(s, 4.35, 5.2, 3.9, 0.3, "Rows = actual K-means segment", size=11, color=MUTED, italic=True)
box(s, 0.74, 4.85, 3.2, 1.1, "Why not just re-run K-means? The classifier gives a fast, fixed scoring rule "
	"that CRM tools can apply daily.", size=12, color=MUTED)
box(s, 9.2, 1.45, 3.4, 0.35, "LABEL AGREEMENT", size=11, color=ON_NAVY_MUTED, bold=True)
box(s, 9.2, 1.8, 3.4, 0.9, f"{acc * 100:.1f}%", size=48, color=WHITE, bold=True)
box(s, 9.2, 2.75, 3.3, 0.7, f"{correct} of {test_n} held-out customers placed in the same segment as K-means.",
	size=13, color=ON_NAVY_MUTED)
box(s, 9.2, 3.75, 3.3, 0.35, "READ WITH CARE", size=11, color=RGBColor(0xFF, 0xB4, 0x9C), bold=True)
box(s, 9.2, 4.1, 3.3, 1.7, "Labels were built from the same RFM features: this measures "
	"reproducibility, not future purchases (FAQ 6, 7).", size=13, color=WHITE)

# 12. Recommendations
s = new_slide(L_TITLE_ONLY, "Recommendations: protect the 39% first")
cols = [
	("High-Value Active", "Protect & grow", BLUE, [
		"Loyalty tier: early access, volume pricing, a named contact for wholesale buyers",
		"Early-warning trigger when a customer passes ~60 days without an order",
		"Cross-sell from their own purchase history ahead of the Q4 peak",
	]),
	("Low-Value Lapsing", "Win back cheaply", GREY, [
		"Automated reactivation emails at 90 / 120 / 180 days since last order",
		"Entry offer on best-selling products to earn the second order",
		"Keep out of costly channels until they re-engage",
	]),
]
for i, (name, verb, col, items) in enumerate(cols):
	x = 0.74 + i * 4.05
	rect(s, x, 1.55, 3.85, 3.4, fill=ICE)
	rect(s, x + 0.25, 1.8, 0.24, 0.24, fill=col, shape=MSO_SHAPE.OVAL)
	box(s, x + 0.6, 1.72, 3.1, 0.4, name, size=15, color=NAVY, bold=True)
	box(s, x + 0.25, 2.15, 3.4, 0.4, verb, size=20, color=NAVY, bold=True)
	bullets(s, x + 0.25, 2.7, 3.4, 2.2, items, size=13, gap=6)
rect(s, 8.95, 1.55, 3.65, 1.62, fill=NAVY)
box(s, 9.2, 1.68, 3.2, 0.3, "AT RISK", size=11, color=RGBColor(0xFF, 0xB4, 0x9C), bold=True)
box(s, 9.2, 1.95, 3.2, 0.6, f"≈ {gbp(risk_loss, k=True)}", size=32, color=WHITE, bold=True)
box(s, 9.2, 2.6, 3.25, 0.55, f"if 5% of High-Value Active (≈{risk_n:.0f}) stop buying", size=12, color=ON_NAVY_MUTED)
rect(s, 8.95, 3.33, 3.65, 1.62, fill=ICE)
box(s, 9.2, 3.46, 3.2, 0.3, "WIN-BACK UPSIDE", size=11, color=GAIN, bold=True)
box(s, 9.2, 3.73, 3.2, 0.6, f"≈ {gbp(winback_gain, k=True)}", size=32, color=NAVY, bold=True)
box(s, 9.2, 4.38, 3.25, 0.55, f"if 10% of Lapsing (≈{winback_n}) place one more order at £{lvl_aov:.0f}", size=12, color=MUTED)
box(s, 0.74, 5.12, 11.85, 0.85, [
	[("Test before scaling: ", {"bold": True, "color": NAVY}),
	 ("run each action as an A/B test against a held-out control group and measure incremental revenue.", {})],
	[(f"Illustrative sizing from segment averages: at-risk = 5% × {hva['Size']:,} × {gbp(hva['Monetary'])}; "
	  f"upside = 10% × {lvl['Size']:,} × £{lvl_aov:.0f} average order (£{lvl['Monetary']:.0f} ÷ {lvl['Frequency']:.2f} orders).",
	  {"size": 11, "color": MUTED, "italic": True})]], size=13)

# 13. Dashboard
s = new_slide(L_TITLE_ONLY, "Interactive dashboard deliverable")
pic = s.shapes.add_picture(str(DASH_PNG), Inches(0.74), Inches(1.5), height=Inches(4.45))
pic.line.color.rgb = LINE
pic.line.width = Pt(1)
px = 0.74 + pic.width / 914400 + 0.4
views = [("Overview", "KPIs, monthly sales, spend distribution, countries"),
		 ("Segments", "Cluster sizes and RFM profile by segment"),
		 ("Model review", "k selection, confusion matrix, label accuracy"),
		 ("Method & notes", "Pipeline, caveats, how to read the scores")]
for i, (h, t) in enumerate(views):
	y = 1.55 + i * 0.95
	badge(s, px, y + 0.05, 0.45, str(i + 1), fill=BLUE, size=13)
	box(s, px + 0.65, y, 12.6 - px - 0.65, 0.35, h, size=15, color=NAVY, bold=True)
	box(s, px + 0.65, y + 0.34, 12.6 - px - 0.65, 0.55, t, size=12, color=MUTED)
box(s, px, 5.4, 12.6 - px, 0.55, "Reads the pipeline's JSON output: re-run main.py and refresh.", size=12,
	color=MUTED, italic=True)

# 14. Limitations & next steps
s = new_slide(L_TITLE_ONLY, "Limitations and next steps")
lim = ["Retrospective: describes past behaviour; it does not forecast it",
	   "Two segments are coarse for campaign planning",
	   "Classifier reproduces K-means labels (circular target)",
	   "One year of data, ~89% UK; wholesale buyers mixed with consumers",
	   "RFM only: no product, basket or channel features"]
nxt = ["Time-split repurchase model: RFM before a cutoff, predict the next 90 days (FAQ 7)",
	   "Offer k = 4 as an operational sub-segmentation (FAQ 3)",
	   "Validate segment names and actions with marketing stakeholders",
	   "Add category mix, basket size, tenure and returns features",
	   "Persist scaler + model and A/B-test the campaigns"]
rect(s, 0.74, 1.55, 5.5, 4.4, fill=ICE)
box(s, 1.0, 1.75, 5.0, 0.4, "Limitations", size=18, color=NAVY, bold=True)
bullets(s, 1.0, 2.3, 5.0, 3.5, lim, size=16, gap=10, bullet_color=GREY)
arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.38), Inches(3.45), Inches(0.5), Inches(0.5))
arr.fill.solid()
arr.fill.fore_color.rgb = BLUE
arr.line.fill.background()
rect(s, 7.05, 1.55, 5.55, 4.4, fill=NAVY)
box(s, 7.3, 1.75, 5.0, 0.4, "Next steps", size=18, color=WHITE, bold=True)
bullets(s, 7.3, 2.3, 5.05, 3.5, nxt, size=16, gap=10, color=WHITE, bullet_color=SKY)

# 15/16: Questions + Thank You are template slides 7 and 8 (kept as-is)
s_q, s_ty = prs.slides[6], prs.slides[7]

# 17. Appendix divider
s_app = new_slide(L_SECTION, "Appendix: anticipated questions")
for ph in s_app.shapes:
	pass
box(s_app, 0.8, 5.05, 11.5, 0.5, "Backup slides for methodology deep-dives · FAQ 1–9", size=16, color=MUTED)

# 18. FAQ index
s_idx = new_slide(L_TITLE_ONLY, "FAQ index")
faq_titles = [
	"Why log-transform and then standardise RFM?",
	"Why only two segments?",
	"What would k = 4 look like?",
	"Why did DBSCAN and hierarchical do worse?",
	"What did cleaning remove, step by step?",
	"Isn't 99% accuracy a sign of leakage?",
	"How would you build a true future-purchase model?",
	"How stable are the clusters?",
	"Why only RFM, not product or country features?",
]
idx_cards = []
for i, q in enumerate(faq_titles):
	r_, c_ = divmod(i, 3)
	x, y = 0.74 + c_ * 4.0, 1.55 + r_ * 1.45
	card = rect(s_idx, x, y, 3.8, 1.25, fill=ICE)
	b = badge(s_idx, x + 0.2, y + 0.35, 0.55, str(i + 1), fill=NAVY, size=15)
	t = box(s_idx, x + 0.95, y + 0.1, 2.7, 1.05, q, size=13, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
	idx_cards.append((card, b, t))

faq_slides = []


def faq(n, title):
	s = new_slide(L_ONE_THIRD, title)
	for r in s.shapes.title.text_frame.paragraphs[0].runs:
		r.font.size = Pt(26)
	box(s, 0.95, 1.5, 2.5, 0.4, f"FAQ {n}", size=14, color=ON_NAVY_MUTED, bold=True)
	back = box(s, 0.95, 5.35, 2.4, 0.35, "↩ FAQ index", size=12, color=ON_NAVY_MUTED)
	back.click_action.target_slide = s_idx
	faq_slides.append(s)
	return s


def answer(s, text, y=0.75, h=0.9, size=16):
	box(s, 4.9, y, 7.7, h, text, size=size, color=NAVY, bold=True)


# FAQ 1
s = faq(1, "Why log-transform and then standardise RFM?")
answer(s, "K-means uses Euclidean distance, so raw RFM would be dominated by a few huge spenders and by Monetary's scale.")
cd = CategoryChartData()
cd.categories = ["Recency", "Frequency", "Monetary"]
cd.add_series("Raw", [1.24, 11.84, 20.35])
cd.add_series("After log1p", [-0.38, 1.21, 0.36])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.9), Inches(1.75), Inches(4.3), Inches(3.9), cd)
ch = gf.chart
chart_title(ch, "Skewness: raw vs log1p")
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
style_axes(ch, grid=None, show_val_axis=False)
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(11)
ch.plots[0].gap_width = 60
labels_on(ch.plots[0], "0.0", size=10)
ch.plots[0].series[0].format.fill.solid()
ch.plots[0].series[0].format.fill.fore_color.rgb = GREY
ch.plots[0].series[1].format.fill.solid()
ch.plots[0].series[1].format.fill.fore_color.rgb = BLUE
bullets(s, 9.45, 1.85, 3.15, 3.9, [
	[("log1p ", {"bold": True}), ("compresses the long tail (max spend £265k vs median £614); log1p is safe at 0.", {})],
	[("StandardScaler ", {"bold": True}), ("gives R, F and M mean 0 and SD 1, so each feature weighs equally.", {})],
	[("Order matters: ", {"bold": True}), ("scaling first would keep the skew.", {})]], size=13, gap=8)

# FAQ 2
s = faq(2, "Why only two segments?")
answer(s, "Because the data's own structure says so: k = 2 has clearly the best silhouette, and the split is very stable.")
xy = XyChartData()
for cl, name, col in [(0, "High-Value Active", BLUE), (1, "Low-Value Lapsing", GREY)]:
	ser = xy.add_series(name)
	for p in dash["cluster_points"]:
		if p["cluster"] == cl:
			ser.add_data_point(p["x"], p["y"])
gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(4.9), Inches(1.7), Inches(4.6), Inches(4.1), xy)
ch = gf.chart
chart_title(ch, "Customers in scaled R–F space (1,200 sample)", size=12)
style_axes(ch, grid=LINE, size=10)
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(10)
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
ch.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
axis_title(ch.category_axis, "Recency (z)", size=10)
axis_title(ch.value_axis, "Frequency (z)", size=10)
for ser, col in zip(ch.plots[0].series, [BLUE, GREY]):
	ser.format.line.fill.background()
	ser.marker.style = XL_MARKER_STYLE.CIRCLE
	ser.marker.size = 4
	ser.marker.format.fill.solid()
	ser.marker.format.fill.fore_color.rgb = col
	ser.marker.format.line.fill.background()
bullets(s, 9.75, 1.85, 2.85, 4.0, [
	"Silhouette: 0.43 at k = 2, then 0.34, 0.33 … 0.28 at k = 10",
	"Elbow is smooth, so it doesn't pick a k",
	f"Seed and bootstrap ARI ≈ {stab['bootstrap_ari_mean']:.2f}",
	"Behaviour is a continuum; finer cuts slice the same gradient (see k = 4, FAQ 3)"], size=13, gap=8)

# FAQ 3
s = faq(3, "What would k = 4 look like?")
answer(s, f"Four readable segments at a lower silhouette ({cmp4['KMeans']['silhouette']:.2f} vs 0.43). The extremes stay put; the middle of the gradient gets its own groups.")
alt = sorted(res["alternative_profile"], key=lambda r: -r["RevenuePct"])
rows = [["Segment", "Customers", "Recency", "Orders", "Avg. spend", "Revenue"]]
for r in alt:
	rows.append([r["Segment"], f"{r['Size']:,}", f"{r['Recency']:.0f} d", f"{r['Frequency']:.1f}",
				 gbp(r["Monetary"]), f"{r['RevenuePct']:.0f}%"])
tbl = s.shapes.add_table(len(rows), 6, Inches(4.9), Inches(1.85), Inches(7.7), Inches(2.6)).table
widths = [2.2, 1.1, 1.05, 0.95, 1.25, 1.15]
for c, w in enumerate(widths):
	tbl.columns[c].width = Inches(w)
for r_i, row in enumerate(rows):
	for c_i, val in enumerate(row):
		cell = tbl.cell(r_i, c_i)
		cell.text = val or " "
		cell.vertical_anchor = MSO_ANCHOR.MIDDLE
		cell.margin_left = cell.margin_right = Inches(0.08)
		p = cell.text_frame.paragraphs[0]
		p.alignment = PP_ALIGN.LEFT if c_i == 0 else PP_ALIGN.RIGHT
		run = p.runs[0]
		run.font.name, run.font.size = FONT, Pt(13)
		cell.fill.solid()
		if r_i == 0:
			cell.fill.fore_color.rgb = NAVY
			run.font.color.rgb, run.font.bold = WHITE, True
		else:
			cell.fill.fore_color.rgb = ICE if r_i % 2 else WHITE
			run.font.color.rgb = INK
			run.font.bold = c_i in (0, 5)
bullets(s, 4.9, 4.7, 7.7, 1.3, [
	"All 709 top customers sit inside High-Value Active (k = 2), and all 1,571 lapsed customers inside Low-Value Lapsing",
	"The two Mid-Value groups straddle the k = 2 line: Cooling is the retention watch-list; Active are recent buyers to convert to a second order"], size=13, gap=6)

# FAQ 4
s = faq(4, "Why did DBSCAN and hierarchical do worse?")
answer(s, "The data is one continuous cloud, not dense islands. That suits centroid methods and defeats density methods.")
rows = [["", "K-means", "Hierarchical", "DBSCAN"],
		["Silhouette (k = 2)", f"{cmp2['KMeans']['silhouette']:.3f}", f"{cmp2['Hierarchical']['silhouette']:.3f}", f"{cmp2['DBSCAN']['silhouette']:.3f}"],
		["Davies-Bouldin (k = 2)", f"{cmp2['KMeans']['davies_bouldin']:.2f}", f"{cmp2['Hierarchical']['davies_bouldin']:.2f}", f"{cmp2['DBSCAN']['davies_bouldin']:.2f}"],
		["Silhouette (k = 4)", f"{cmp4['KMeans']['silhouette']:.3f}", f"{cmp4['Hierarchical']['silhouette']:.3f}", "n/a"],
		["Noise points", "0", "0", f"{cmp2['DBSCAN']['noise_points']}"]]
tbl = s.shapes.add_table(len(rows), 4, Inches(4.9), Inches(1.85), Inches(7.7), Inches(2.1)).table
for c, w in enumerate([2.9, 1.6, 1.6, 1.6]):
	tbl.columns[c].width = Inches(w)
for r_i, row in enumerate(rows):
	for c_i, val in enumerate(row):
		cell = tbl.cell(r_i, c_i)
		cell.text = val or " "
		cell.vertical_anchor = MSO_ANCHOR.MIDDLE
		p = cell.text_frame.paragraphs[0]
		p.alignment = PP_ALIGN.LEFT if c_i == 0 else PP_ALIGN.CENTER
		run = p.runs[0]
		run.font.name, run.font.size = FONT, Pt(13)
		cell.fill.solid()
		if r_i == 0:
			cell.fill.fore_color.rgb = NAVY
			run.font.color.rgb, run.font.bold = WHITE, True
		else:
			cell.fill.fore_color.rgb = ICE if r_i % 2 else WHITE
			run.font.color.rgb = NAVY if c_i == 1 else INK
			run.font.bold = c_i <= 1
bullets(s, 4.9, 4.2, 3.7, 1.8, [
	[("DBSCAN ", {"bold": True}), ("(eps 0.5, min_samples 5) finds one dominant dense region plus noise; one eps can't fit both dense and sparse areas.", {})]],
	size=13)
bullets(s, 8.85, 4.2, 3.75, 1.8, [
	[("Ward hierarchical ", {"bold": True}), ("is close at k = 2 but merges greedily and can't undo early merges; it also has no predict step.", {})]],
	size=13)

# FAQ 5
s = faq(5, "What did cleaning remove, step by step?")
answer(s, f"Five rule-based steps took {sig['raw_rows']:,} rows to {sig['clean_rows']:,}. Each rule has a business reason.")
reasons = ["RFM is per customer; anonymous rows can't be attributed",
		   "Invoices starting 'C' are returns, not purchases",
		   "Prices below £0.01 are adjustments or errors",
		   "Drops service codes (POST, BANK CHARGES) but also lettered variants like 85123A",
		   "Exact repeats would double-count spend"]
rows = [["Step", "Rows removed", "Remaining", "Why"]]
for st, why in zip(steps, reasons):
	rows.append([st["step"], f"{st['removed']:,}", f"{st['remaining']:,}", why])
tbl = s.shapes.add_table(len(rows), 4, Inches(4.9), Inches(1.85), Inches(7.7), Inches(3.4)).table
for c, w in enumerate([2.0, 1.25, 1.25, 3.2]):
	tbl.columns[c].width = Inches(w)
for r_i, row in enumerate(rows):
	for c_i, val in enumerate(row):
		cell = tbl.cell(r_i, c_i)
		cell.text = val or " "
		cell.vertical_anchor = MSO_ANCHOR.MIDDLE
		cell.margin_left = cell.margin_right = Inches(0.08)
		p = cell.text_frame.paragraphs[0]
		p.alignment = PP_ALIGN.RIGHT if c_i in (1, 2) else PP_ALIGN.LEFT
		run = p.runs[0]
		run.font.name, run.font.size = FONT, Pt(12)
		cell.fill.solid()
		if r_i == 0:
			cell.fill.fore_color.rgb = NAVY
			run.font.color.rgb, run.font.bold = WHITE, True
		else:
			cell.fill.fore_color.rgb = ICE if r_i % 2 else WHITE
			run.font.color.rgb = INK
			run.font.bold = c_i == 0
box(s, 4.9, 5.45, 7.7, 0.5, "Every step prints its count at run time and is saved to outputs/analysis_results.json.",
	size=12, color=MUTED, italic=True)

# FAQ 6
s = faq(6, "Isn't 99% accuracy a sign of leakage?")
answer(s, "It's circular by design, and we say so. The score measures reproducibility of the K-means labels, not forecasting.")
for i, (h, t, col) in enumerate([
		("What it is", "The target (segment) was created by K-means from the same scaled R, F, M. K-means boundaries are "
		 "simple straight lines in that space, so a Random Forest recovers them almost perfectly.", NAVY),
		("Why it's still useful", "It gives a fixed, fast scoring rule for new customers, and it confirms the segments are "
		 f"learnable and consistent. The {test_n - correct} errors sit near the boundary.", BLUE),
		("What it is NOT", "Evidence that we can predict who will buy next. That needs a target from a later time window "
		 "(FAQ 7).", RISK)]):
	y = 1.85 + i * 1.35
	rect(s, 4.9, y, 7.7, 1.2, fill=ICE)
	box(s, 5.15, y + 0.12, 7.2, 0.35, h, size=15, color=col, bold=True)
	box(s, 5.15, y + 0.47, 7.25, 0.7, t, size=13, color=INK)

# FAQ 7
s = faq(7, "How would you build a true future-purchase model?")
answer(s, "Split by time: build features only from the past, and define the target in a later window the model never saw.")
tl_y = 2.3
segs_tl = [("Feature window", "1 Dec 2010 – 31 Aug 2011", "Compute RFM (+ extra features)", 4.9, 4.4, NAVY),
		   ("Outcome window", "1 Sep – 30 Nov 2011", "Target: purchased again in 90 days?", 9.35, 3.25, BLUE)]
for name, dates, what, x, w, col in segs_tl:
	rect(s, x, tl_y, w, 0.75, fill=col)
	box(s, x + 0.15, tl_y, w - 0.3, 0.75, [[(name, {"bold": True, "size": 14})], [(dates, {"size": 12})]],
		size=13, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
	box(s, x + 0.15, tl_y + 0.85, w - 0.3, 0.5, what, size=12, color=MUTED)
box(s, 8.95, tl_y - 0.42, 1.0, 0.35, "cutoff", size=11, color=RISK, bold=True, align=PP_ALIGN.CENTER)
ln = s.shapes.add_connector(1, Inches(9.3), Inches(tl_y - 0.1), Inches(9.3), Inches(tl_y + 0.85))
ln.line.color.rgb = RISK
ln.line.width = Pt(2)
bullets(s, 4.9, 3.85, 7.7, 2.1, [
	[("Validate forward in time: ", {"bold": True}), ("train on one cutoff, test on a later one; never shuffle across time.", {})],
	[("Metrics: ", {"bold": True}), ("ROC-AUC, precision-recall AUC, and lift in the top decile, compared against a simple 'most recent buyers' baseline.", {})],
	[("Use: ", {"bold": True}), ("rank customers by repurchase probability to target retention spend.", {})]], size=13, gap=6)

# FAQ 8
s = faq(8, "How stable are the clusters?")
answer(s, "Very stable: refitting with different seeds or on 80% subsamples reproduces almost exactly the same split.")
cd = CategoryChartData()
cd.categories = ["Seeds (20 × single init)", "Bootstrap (20 × 80% sample)"]
cd.add_series("Mean ARI", [stab["seed_ari_mean"], stab["bootstrap_ari_mean"]])
cd.add_series("Minimum ARI", [stab["seed_ari_min"], stab["bootstrap_ari_min"]])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.9), Inches(1.75), Inches(4.6), Inches(4.1), cd)
ch = gf.chart
chart_title(ch, "Adjusted Rand index vs full-data labels")
style_axes(ch, grid=None, show_val_axis=False)
ch.value_axis.minimum_scale = 0
ch.value_axis.maximum_scale = 1.1
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(11)
ch.plots[0].gap_width = 70
labels_on(ch.plots[0], "0.000", size=11)
ch.plots[0].series[0].format.fill.solid()
ch.plots[0].series[0].format.fill.fore_color.rgb = NAVY
ch.plots[0].series[1].format.fill.solid()
ch.plots[0].series[1].format.fill.fore_color.rgb = SKY
bullets(s, 9.75, 1.85, 2.85, 4.0, [
	"ARI = 1 means identical grouping; 0 means chance",
	"Seed test uses a single initialisation, a harder test than production (n_init = 10)",
	"Bootstrap compares labels on the sampled customers only",
	f"Worst case across 40 refits: {min(stab['seed_ari_min'], stab['bootstrap_ari_min']):.3f}"], size=13, gap=8)

# FAQ 9
s = faq(9, "Why only RFM, not product or country features?")
answer(s, "RFM is what the brief asked us to analyse, it exists for every customer, and it stays explainable to marketers.")
feats3 = ["Recency", "Frequency", "Monetary"]
gx, gy, cs = 5.9, 2.2, 1.0
box(s, 4.9, 1.8, 4.2, 0.35, "Pearson correlation (raw RFM)", size=13, color=NAVY, bold=True)
for i, a in enumerate(feats3):
	box(s, gx - 1.05, gy + i * cs, 1.0, cs, a, size=11, color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
	box(s, gx + i * cs, gy + 3 * cs + 0.05, cs, 0.3, a, size=11, color=MUTED, align=PP_ALIGN.CENTER)
	for j, b in enumerate(feats3):
		v = dash["correlations"][a][b]
		if i == j:
			fill, fc = NAVY, WHITE
		elif abs(v) >= 0.4:
			fill, fc = BLUE, WHITE
		elif abs(v) >= 0.2:
			fill, fc = SKY, NAVY
		else:
			fill, fc = ICE, NAVY
		rect(s, gx + j * cs + 0.03, gy + i * cs + 0.03, cs - 0.06, cs - 0.06, fill=fill, radius=0.05)
		box(s, gx + j * cs, gy + i * cs, cs, cs, f"{v:.2f}", size=14, color=fc, bold=True,
			align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, 9.35, 1.85, 3.25, 4.1, [
	"F and M overlap (0.53) but aren't redundant; R adds an independent signal",
	"More features make distances less meaningful and segments harder to name",
	"Country is ~89% UK: little variation to segment on",
	"Next: category mix, basket size, tenure, return rate, possibly with PCA"], size=13, gap=8)

# ---------- index hyperlinks ----------
for shapes, fs in zip(idx_cards, faq_slides):
	for shp in shapes:
		shp.click_action.target_slide = fs

# ---------- remove template sample slides 3–6 and reorder ----------
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
for sid in ids[2:6]:
	prs.part.drop_rel(sid.rId)
	sldIdLst.remove(sid)
ids = list(sldIdLst)  # [title, agenda, questions, thankyou, main 3..14, appendix, index, faq1..9]
title_, agenda_, q_, ty_ = ids[:4]
rest = ids[4:]
main, tail = rest[:12], rest[12:]
for sid in list(sldIdLst):
	sldIdLst.remove(sid)
for sid in [title_, agenda_] + main + [q_, ty_] + tail:
	sldIdLst.append(sid)

# ---------- presenter notes (final slide order) ----------
# Main talk: verbatim opening, numeric talking points, verbatim transition.
# FAQs: one-line answer, 30-second explanation, follow-up, and what not to claim.


def talk(opening, points, transition=None):
	lines = [f"OPEN: “{opening}”"] + [f"• {p}" for p in points]
	if transition:
		lines.append(f"NEXT: “{transition}”")
	return "\n".join(lines)


def faq_note(one_line, explain, follow_ups, dont=None):
	lines = [f"1 · ONE-LINE ANSWER (say first): “{one_line}”", "", f"2 · IF THEY WANT MORE (~30s): {explain}", "",
			 "3 · IF THEY PUSH FURTHER:"]
	lines += [f"• Q: {q}\n   A: “{a}”" for q, a in follow_ups]
	if dont:
		lines += ["", f"DON'T SAY: {dont}"]
	return "\n".join(lines)


slide_no = {id(s._element): i + 1 for i, s in enumerate(prs.slides)}
faq_no = {n: slide_no[id(fs._element)] for n, fs in enumerate(faq_slides, 1)}
idx_no = slide_no[id(s_idx._element)]
q_no = slide_no[id(s_q._element)]
lost = test_n - correct

SCRIPT = [
	# 1 Title
	(15, talk("Good [morning/afternoon]. I'm [name], and my capstone asks one question: who actually drives the "
			  "revenue in an online retail business?",
			  ["Fill in [Presenter name] and [Date] on the slide before presenting."],
			  "Here's how the next ten minutes are organised.")),
	# 2 Agenda
	(20, talk("I'll move through five parts, from the business problem to what the business should do next.",
			  ["Problem & data → customer behaviour → building segments → segments & drivers → acting on it.",
			   "Methodology deep dives are in an FAQ appendix after the close; invite the panel to use it in Q&A."],
			  "Let me start with the problem.")),
	# 3 Business problem
	(40, talk("An e-commerce company wants to spend its marketing and retention budget where it matters, but it "
			  "doesn't yet know which customers matter.",
			  ["Four objectives straight from the brief: segment, explain, operationalise, act.",
			   "Segment means comparing K-means, hierarchical clustering and DBSCAN on purchasing behaviour.",
			   "Success is an actionable segmentation, not just a high metric."],
			  "To answer that, I needed clean, customer-level data.")),
	# 4 Data
	(35, talk("I used the public UCI Online Retail dataset: one year of transactions from a UK gift retailer.",
			  [f"{sig['raw_rows']:,} rows → {sig['clean_rows']:,} after five cleaning rules; {n_cust:,} customers.",
			   f"Biggest cut: {steps[0]['removed']:,} rows with no CustomerID, unavoidable for customer-level analysis.",
			   f"Step-by-step detail is FAQ 5 (slide {faq_no[5]})."],
			  "Before modelling, I looked at what the transactions tell us.")),
	# 5 EDA
	(45, talk("Three patterns shaped everything that followed.",
			  [f"Seasonality: revenue climbs from September to a {gbp(months['revenue'][peak_i])} peak in November: gift buying.",
			   f"Concentration: about {uk_share:.0f}% of transactions ({uk_rev_share:.0f}% of revenue) are from the UK.",
			   "Skew: the median customer spends £614, the mean is £1,853, the largest about £265k.",
			   "If asked: December 2011 is low only because the data stops on 9 December."],
			  "That skew is exactly why I didn't cluster the raw numbers.")),
	# 6 RFM
	(40, talk("I summarised each customer with three behavioural features: Recency, Frequency and Monetary value.",
			  ["Medians: 51 days since last purchase, 2 orders, £614 spend.",
			   "log1p then StandardScaler: Monetary skewness falls from 20.4 to 0.4, so no single feature or big "
			   "spender dominates the distance.",
			   f"Why this order: FAQ 1 (slide {faq_no[1]})."],
			  "With comparable features, the next question was how many segments.")),
	# 7 Choosing k
	(55, talk("I let the data choose the number of segments rather than picking one that sounded good.",
			  ["K-means for k = 2 to 10, n_init 10, fixed seed.",
			   "Silhouette peaks at k = 2 with 0.43, then drops to 0.34 at k = 3 and keeps declining.",
			   "The elbow curve is smooth with no sharp bend, so the silhouette decides.",
			   f"I keep k = 4 as a finer operational view: FAQ 3 (slide {faq_no[3]})."],
			  "Next, I checked whether K-means was the right algorithm at all.")),
	# 8 Algorithm comparison
	(55, talk("The brief asked me to compare three families of clustering, so I ran all three on the same features.",
			  [f"Silhouette: K-means {cmp2['KMeans']['silhouette']:.3f}, Ward hierarchical {cmp2['Hierarchical']['silhouette']:.3f}, "
			   f"DBSCAN {cmp2['DBSCAN']['silhouette']:.3f}.",
			   f"Davies-Bouldin agrees: K-means is lowest at {cmp2['KMeans']['davies_bouldin']:.2f}.",
			   f"DBSCAN finds one dominant dense region plus {cmp2['DBSCAN']['noise_points']} noise points: behaviour is a "
			   "continuum, not separate islands.",
			   "Practical tie-breaker: only K-means can score a new customer without refitting."],
			  "So what do the two K-means segments actually look like?")),
	# 9 Segments
	(70, talk("Two segments, and they are very unequal.",
			  [f"High-Value Active: {hva['Size']:,} customers, last order about {hva['Recency']:.0f} days ago, about "
			   f"{hva['Frequency']:.0f} orders, about {gbp(hva['Monetary'])} each.",
			   f"Low-Value Lapsing: {lvl['Size']:,} customers, about {lvl['Recency']:.0f} days since last order, 1–2 orders, "
			   f"about {gbp(lvl['Monetary'])}.",
			   f"Headline: {hva_share:.0f}% of customers generate {hva['RevenuePct']:.0f}% of revenue, "
			   f"{gbp(hva_rev)} of {gbp(total_rev)}. Pause here.",
			   f"Names come from each cluster's centroid (value tier + activity status), not chosen by hand; "
			   f"the split is stable, ARI ≈ {stab['bootstrap_ari_mean']:.2f} (FAQ 8)."],
			  "The natural follow-up is: what separates these two groups?")),
	# 10 Drivers
	(45, talk("I looked at that from two independent angles, and they agree.",
			  [f"Centroid gap in standard units: Frequency {gap['Frequency']:.2f}, Monetary {gap['Monetary']:.2f}, "
			   f"Recency {gap['Recency']:.2f}.",
			   f"Random Forest importance: Frequency {fi['Frequency']:.0%}, Monetary {fi['Monetary']:.0%}, "
			   f"Recency {fi['Recency']:.0%}.",
			   "Business translation: earning the second and third order is the lever that moves a customer up."],
			  "That Random Forest also solves a practical problem: scoring new customers.")),
	# 11 Scoring
	(50, talk("Once segments exist, the business needs to assign tomorrow's customers to them.",
			  [f"200-tree Random Forest, stratified 80/20 split: {correct} of {test_n} held-out customers match K-means, "
			   f"{acc * 100:.1f}%.",
			   "Say this before the panel does: the labels came from K-means on the same features, so this measures "
			   f"reproducibility, not future purchases (FAQ 6, 7 on slides {faq_no[6]}–{faq_no[7]}).",
			   "If asked: accuracy varies slightly between runs, 99.4–99.7%."],
			  "So how should the business act on these segments?")),
	# 12 Recommendations
	(70, talk(f"My core recommendation: the business should protect the {hva_share:.0f}% before chasing the {lvl_share:.0f}%.",
			  [f"Risk: losing just 5% of High-Value Active, about {risk_n:.0f} customers, puts about {gbp(risk_loss, k=True)} at risk.",
			   f"Upside: winning back 10% of Lapsing, about {winback_n} customers, for one more order adds about "
			   f"{gbp(winback_gain, k=True)}.",
			   "Actions: a loyalty tier and a ~60-day no-order trigger for the top segment; low-cost automated "
			   "win-back for the rest.",
			   "These are illustrative sizes from segment averages; every action should be A/B-tested against a "
			   "control group."],
			  "All of this is available to the business in an interactive dashboard.")),
	# 13 Dashboard
	(25, talk("The dashboard is the interactive deliverable from the brief.",
			  ["Four views: Overview, Segments, Model review, Method & notes. It reads the pipeline's JSON, so a "
			   "rerun refreshes it.",
			   "OPTIONAL LIVE DEMO (~60s): only if the clock shows 8:30 or less on arriving here.",
			   "   Before the talk: from CapstoneEcommerceSegmentation run  python -m http.server 8000  and open "
			   "http://localhost:8000/dashboards/index.html in a browser tab.",
			   "   Click path: Overview KPIs → Segments → Model review (confusion matrix) → back to the slides.",
			   "   If anything fails: “I'll leave you with the screenshot” and move on."],
			  "Finally, what this work doesn't do yet, and what I'd do next.")),
	# 14 Limitations
	(35, talk("I want to be clear about the limits of this work.",
			  ["Retrospective; coarse at k = 2; circular classifier target; one UK-heavy year; RFM only.",
			   f"The most valuable next step is a time-split repurchase model: FAQ 7 (slide {faq_no[7]}).",
			   f"CLOSE, word for word: “To sum up: {hva_share:.0f}% of customers drive {hva['RevenuePct']:.0f}% of revenue, "
			   "so protecting them is worth far more than chasing the rest. Thank you. I'm happy to take questions.”"])),
	# 15 Questions ?
	(0, "\n".join([
		"Q&A HABITS",
		"1. Repeat or rephrase the question so the whole panel hears it.",
		"2. Give the one-sentence answer first; stop and let them ask for more.",
		"3. Offer evidence: “I have a backup slide on that. Shall I show it?”",
		f"4. Jump by typing the slide number and pressing Enter; come back with {q_no} + Enter.",
		"",
		"ROUTING CHEAT SHEET (question → slide)",
		f"• Why log / why scale / why these features? → FAQ 1 · slide {faq_no[1]}",
		f"• Why only 2 segments? Isn't that too few? → FAQ 2 · slide {faq_no[2]}, then FAQ 3 · slide {faq_no[3]}",
		f"• Why K-means? Why not DBSCAN, hierarchical or GMM? → slide 8, then FAQ 4 · slide {faq_no[4]}",
		f"• How did you clean the data? What about returns? → FAQ 5 · slide {faq_no[5]}",
		f"• 99% accuracy: leakage? overfitting? → FAQ 6 · slide {faq_no[6]}",
		f"• Can you predict churn or future purchases? → FAQ 7 · slide {faq_no[7]}",
		f"• Are the clusters stable or reproducible? → FAQ 8 · slide {faq_no[8]}",
		f"• Why not product / country features? PCA? → FAQ 9 · slide {faq_no[9]}",
		"• How were the £340k / £72k figures calculated? → slide 12 footnote",
		"• How were the segment names decided? → slide 9: centroid value tier + activity status",
		"• Can we see the dashboard? → live demo, or slide 13",
		"",
		"UNPREPARED QUESTION FALLBACK",
		"“That's a good question I haven't tested directly. My expectation is …, and the way I'd check it is ….” "
		"Never guess a number."])),
	# 16 Thank You!
	(0, "IF ASKED TO SUM UP (~20s, word for word):\n"
		f"“Using RFM features and K-means, I found two stable customer segments. {hva_share:.0f}% of customers generate "
		f"{hva['RevenuePct']:.0f}% of revenue, so the business case says: protect that group first, win back the rest "
		"cheaply, and test every action against a control group.”"),
	# 17 Appendix divider
	(0, f"Only enter the appendix when asked. Go straight to the relevant FAQ by typing its slide number + Enter; "
		f"the FAQ index is slide {idx_no}."),
	# 18 FAQ index
	(0, "NAVIGATION\n"
		"• Click a card to jump to that FAQ.\n"
		f"• On any FAQ slide, click “↩ FAQ index” (bottom left of the navy panel), or type {idx_no} + Enter.\n"
		f"• FAQ n is slide n + {faq_no[1] - 1}  (FAQ 1 = slide {faq_no[1]}, FAQ 9 = slide {faq_no[9]}).\n"
		f"• Back to the Questions slide: {q_no} + Enter."),
	# FAQ 1
	(0, faq_note(
		"Because K-means uses distance, and without the log and scaling a handful of big spenders and the pound "
		"scale would decide the clusters.",
		"Raw skewness is 20.4 for Monetary and 11.8 for Frequency; after log1p it's 0.4 and 1.2. StandardScaler then "
		"gives each feature mean 0 and standard deviation 1, so all three weigh equally. The log comes first because "
		"scaling alone doesn't change the shape of a distribution.",
		[("Why not remove outliers or use RobustScaler?",
		  "The outliers are my most valuable wholesale customers; removing them would hide the segment that matters "
		  "most. The log tames them while keeping them. RobustScaler would be a reasonable sensitivity check.")],
		"that the log makes the data normal: Frequency is still skewed (1.2).")),
	# FAQ 2
	(0, faq_note(
		"Because the silhouette clearly peaks at two, and I didn't want to overrule my own metric for a nicer story.",
		"I tested k from 2 to 10: 0.43 at k = 2, 0.34 at k = 3, declining to 0.28 at k = 10. The elbow is smooth. "
		f"Davies-Bouldin also favours k = 2 ({cmp2['KMeans']['davies_bouldin']:.2f} vs {cmp4['KMeans']['davies_bouldin']:.2f} at k = 4), "
		f"and the split is stable (ARI ≈ {stab['bootstrap_ari_mean']:.2f}). For finer targeting there's k = 4 on the next slide.",
		[("Isn't k = 2 usually highest on silhouette anyway?",
		  "Often, yes. That's why I also checked Davies-Bouldin and stability, and present k = 4 as the operational option.")],
		"that two is the 'true' number of customer types.")),
	# FAQ 3
	(0, faq_note(
		f"Four readable segments with a lower silhouette, {cmp4['KMeans']['silhouette']:.2f}, useful as an operational "
		"sub-segmentation.",
		"High-Value Active 709 customers (about 65% of revenue), Mid-Value Cooling 1,179 (24%), Mid-Value Active 855 "
		"(5%, recent but about 2 orders), Low-Value Lapsing 1,571 (6%). The top 709 and bottom 1,571 sit entirely "
		"inside the k = 2 groups; the two mid groups straddle the k = 2 boundary, which shows behaviour is a gradient. "
		"Mid-Value Cooling is the natural retention watch-list.",
		[("Then why not present k = 4 as the main result?",
		  "Because the evidence favours k = 2. k = 4 is a business refinement, and I'd validate it with stakeholders "
		  "and a stability check first.")],
		"that k = 4 nests neatly inside k = 2, or that its stability was tested.")),
	# FAQ 4
	(0, faq_note(
		"Customer behaviour is one continuous cloud, which suits centroid methods and defeats density methods.",
		f"Silhouette {cmp2['KMeans']['silhouette']:.3f} / {cmp2['Hierarchical']['silhouette']:.3f} / "
		f"{cmp2['DBSCAN']['silhouette']:.3f}; Davies-Bouldin {cmp2['KMeans']['davies_bouldin']:.2f} / "
		f"{cmp2['Hierarchical']['davies_bouldin']:.2f} / {cmp2['DBSCAN']['davies_bouldin']:.2f}; at k = 4, K-means "
		f"{cmp4['KMeans']['silhouette']:.3f} vs Ward {cmp4['Hierarchical']['silhouette']:.3f}. DBSCAN with eps 0.5 and "
		f"min_samples 5 finds one dominant region plus {cmp2['DBSCAN']['noise_points']} noise points. Ward is close at "
		"k = 2 but merges greedily, needs O(n²) memory, and can't score new customers.",
		[("Did you tune DBSCAN?",
		  "I used standard defaults. A k-distance plot to choose eps is the right next step, but on a continuous "
		  "cloud tuning mostly trades noise points for one giant cluster."),
		 ("Why not a Gaussian mixture model?",
		  "It's a good extension: soft membership would show how confident each assignment is. It's on my next-steps list.")],
		"that DBSCAN is a bad algorithm. It's the wrong fit for this data shape.")),
	# FAQ 5
	(0, faq_note(
		f"Five transparent rules took {sig['raw_rows']:,} rows to {sig['clean_rows']:,}, and each step is logged.",
		f"Missing CustomerID {steps[0]['removed']:,} (about a quarter), cancellations {steps[1]['removed']:,}, invalid "
		f"prices {steps[2]['removed']:,}, non-numeric stock codes {steps[3]['removed']:,}, exact duplicates "
		f"{steps[4]['removed']:,}.",
		[("Doesn't dropping cancellations overstate spend?",
		  f"Slightly. Netting returns against the original purchases would be more precise; cancellations were "
		  f"{steps[1]['removed'] / sig['raw_rows'] * 100:.1f}% of rows."),
		 ("Why keep only numeric stock codes?",
		  "It removes service codes such as POST, M, D and BANK CHARGES, but it also drops real product variants "
		  "with letter suffixes, like 85123A. That's the cleaning rule I'd revisit first.")],
		f"that all {steps[3]['removed']:,} stock-code rows were non-products.")),
	# FAQ 6
	(0, faq_note(
		"It's circular by design: the score shows the K-means labels are reproducible, not that I can predict "
		"future purchases.",
		"The target was created by K-means from the same scaled R, F and M, and K-means boundaries are straight "
		f"lines in that space, so a Random Forest recovers them almost perfectly: {lost} errors in {test_n}, all near "
		"the boundary. It's a stratified held-out test, so it isn't memorisation. Its value is a fixed, fast scoring "
		"rule for new customers.",
		[("Then why train a classifier at all?",
		  "K-means' own predict could assign new customers too. The Random Forest adds feature importance and lets "
		  "me add non-RFM features later. The honest predictive version is FAQ 7."),
		 ("Why does accuracy change between runs?",
		  "It moves slightly, 99.4 to 99.7%, because a handful of boundary customers flip. It doesn't change any conclusion.")],
		f"“{acc * 100:.1f}% accuracy shows the model predicts customer behaviour.”")),
	# FAQ 7
	(0, faq_note(
		"Split by time: build features only from before a cutoff date and predict purchases in a later window.",
		"Features from 1 Dec 2010 to 31 Aug 2011; target is whether the customer buys between 1 Sep and 30 Nov "
		"2011, a 90-day window. Validate forward in time, never shuffled, and measure ROC-AUC, precision-recall AUC "
		"and top-decile lift against a simple 'most recent buyers' baseline.",
		[("Why haven't you built it?",
		  "About 12 months of data with a strong Q4 season makes one clean train/test pair tight. I prioritised a "
		  "sound segmentation first; it's my top next step."),
		 ("Wouldn't Q4 seasonality bias it?",
		  "Yes. The outcome window includes the peak, so I'd compare against a same-season baseline, or use several "
		  "cutoffs once more data exists.")],
		"that this model has been built or evaluated.")),
	# FAQ 8
	(0, faq_note(
		"Very stable: 40 refits across seeds and subsamples reproduce almost exactly the same split.",
		f"20 seeds with a single initialisation each: mean ARI {stab['seed_ari_mean']:.3f}, minimum "
		f"{stab['seed_ari_min']:.3f}. 20 bootstrap samples of 80%: mean {stab['bootstrap_ari_mean']:.3f}, minimum "
		f"{stab['bootstrap_ari_min']:.3f}. ARI 1 means identical grouping; 0 means chance.",
		[("Is k = 4 as stable?",
		  "I haven't measured it; that's the first check before adopting k = 4 operationally.")],
		"that k = 4 is equally stable.")),
	# FAQ 9
	(0, faq_note(
		"Because the brief asked about recency, frequency and monetary value, they exist for every customer, and "
		"marketers can act on them.",
		"Correlations: Frequency–Monetary 0.53, Recency–Frequency −0.26, Recency–Monetary −0.12, so they're related "
		"but not redundant. More features dilute distance-based clustering and make segments harder to name. "
		f"Country adds little: about {uk_share:.0f}% of transactions are from the UK.",
		[("Would PCA help?",
		  "With three features, no. With a richer set of product and basket features, PCA or feature selection "
		  "before clustering would make sense.")],
		"that RFM captures everything about a customer.")),
]
assert len(SCRIPT) == len(prs.slides), (len(SCRIPT), len(prs.slides))

running = 0
for slide, (secs, text) in zip(prs.slides, SCRIPT):
	running += secs
	head = f"[~{secs}s · running {running // 60}:{running % 60:02d}]\n" if secs else ""
	slide.notes_slide.notes_text_frame.text = head + text

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides), "main talk seconds:", running)
